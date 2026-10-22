# Model finansowy Energypedia.pl – scenariusz bazowy + mnożniki scenariuszy
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

months = []
y, m = 2026, 11
while (y, m) <= (2028, 12):
    months.append((y, m)); m += 1
    if m == 13: y, m = y + 1, 1
N = len(months)
lab = [f"{y}-{m:02d}" for y, m in months]
idx = {l: i for i, l in enumerate(lab)}

def path(points):
    """Interpolacja liniowa między punktami kontrolnymi {miesiąc: wartość}."""
    keys = sorted(points, key=lambda k: idx[k]); out = [0.0] * N
    for i in range(N):
        if i <= idx[keys[0]]: out[i] = points[keys[0]] if i == idx[keys[0]] else 0; continue
        for a, b in zip(keys, keys[1:]):
            ia, ib = idx[a], idx[b]
            if ia <= i <= ib:
                out[i] = points[a] + (points[b] - points[a]) * (i - ia) / (ib - ia); break
        else: out[i] = points[keys[-1]]
    return [round(v, 6) for v in out]

def steps(points):
    """Wartość obowiązuje od wskazanego miesiąca do następnej zmiany."""
    out = [0.0] * N; cur = 0
    for i, l in enumerate(lab):
        if l in points: cur = points[l]
        out[i] = cur
    return out

# ---------- DRIVERS (scenariusz bazowy) ----------
D = {}
D["visits"] = ("Wizyty (sesje) / mies.", path({"2026-11": 0, "2026-12": 25000, "2027-01": 90000, "2027-02": 140000, "2027-03": 200000, "2027-04": 260000, "2027-05": 320000, "2027-06": 400000, "2027-12": 650000, "2028-12": 1100000}))
D["b2b_rate"] = ("Konwersja wizyt -> kwalifikowany lead MŚP", path({"2026-12": 0.0004, "2027-03": 0.0006, "2027-06": 0.0007, "2028-12": 0.0009}))
D["b2b_price"] = ("Przychód za sprzedany lead MŚP [zł]", steps({"2026-12": 160, "2027-04": 190, "2028-01": 210}))
D["b2b_sell"] = ("Odsetek leadów MŚP sprzedanych", steps({"2026-12": 0.6, "2027-03": 0.75, "2027-09": 0.8}))
D["big_deals"] = ("Kontrakty dużych odbiorców (success fee) / mies.", steps({"2027-04": 1, "2027-06": 2, "2027-09": 3, "2028-01": 5}))
D["big_fee"] = ("Średnie wynagrodzenie za kontrakt dużego odbiorcy [zł]", steps({"2027-04": 9000}))
D["hh_rate"] = ("Konwersja wizyt -> lead gospodarstwa domowego (energia)", steps({"2027-01": 0.0025}))
D["hh_price"] = ("Przychód za sprzedany lead GD [zł]", steps({"2027-01": 30}))
D["hh_sell"] = ("Odsetek leadów GD sprzedanych", steps({"2027-01": 0.6}))
D["oze_rate"] = ("Konwersja wizyt -> lead OZE/efektywność (PV, magazyn, pompa, EV)", steps({"2027-02": 0.0012}))
D["oze_price"] = ("Przychód za sprzedany lead OZE [zł]", steps({"2027-02": 110}))
D["oze_sell"] = ("Odsetek leadów OZE sprzedanych", steps({"2027-02": 0.7}))
D["rpm"] = ("Programmatic: przychód na 1000 wizyt [zł]", steps({"2027-01": 10, "2027-07": 12, "2028-01": 14}))
D["spons_art"] = ("Artykuły sponsorowane / native / mies.", steps({"2027-02": 2, "2027-03": 4, "2027-04": 6, "2027-05": 8, "2027-06": 10, "2027-09": 12, "2028-01": 16}))
D["spons_art_price"] = ("Cena artykułu sponsorowanego [zł]", steps({"2027-02": 4500, "2028-01": 6000}))
D["spons_pack"] = ("Aktywne pakiety sponsoringu sekcji/akademii", steps({"2027-03": 1, "2027-05": 2, "2027-06": 3, "2027-10": 4, "2028-01": 6}))
D["spons_pack_price"] = ("Cena pakietu sponsoringu / mies. [zł]", steps({"2027-03": 15000, "2028-01": 18000}))
D["newsletter"] = ("Sponsoring newslettera / mies. [zł]", steps({"2027-03": 3000, "2027-06": 6000, "2027-10": 10000, "2028-01": 15000}))
D["pro_subs"] = ("Abonenci Energypedia PRO (B2B)", path({"2027-04": 10, "2027-06": 40, "2027-12": 150, "2028-12": 450}))
D["pro_price"] = ("Średnia opłata PRO netto / mies. [zł]", steps({"2027-04": 249, "2028-01": 299}))
D["data_rev"] = ("Raporty, licencje danych, API / mies. [zł]", steps({"2027-04": 8000, "2027-06": 20000, "2027-10": 35000, "2028-01": 50000, "2028-07": 70000}))
D["events"] = ("Wydarzenia regionalne / webinary sponsorowane / mies.", steps({"2027-03": 1, "2027-06": 2, "2027-09": 2, "2028-01": 3}))
D["event_margin"] = ("Marża netto na wydarzeniu [zł]", steps({"2027-03": 18000}))
# koszty
# Serwis, asystent AI i treści robimy samodzielnie (naczelny + wirtualny zespół + Claude Code); bez wykonawców, agencji i autorów zewnętrznych.
# Etaty tylko poza redakcją i technologią: growth marketer (XII 2026), key account i weryfikacja leadów (I 2027), drugi weryfikator i koordynator regionów (III 2027).
D["payroll"] = ("Wynagrodzenia zespołu poza redakcją i technologią (z narzutami) [zł]", steps({"2026-12": 14000, "2027-01": 33000, "2027-03": 50000, "2028-01": 62000}))
D["experts"] = ("Rada Ekspertów (bez wynagrodzeń; współpraca honorowa) [zł]", steps({"2026-11": 0}))
D["marketing"] = ("Marketing i PR (płatne, kampanie, influencerzy) [zł]", steps({"2026-12": 40000, "2027-01": 180000, "2027-02": 120000, "2027-03": 110000, "2027-04": 100000, "2027-06": 120000, "2027-07": 80000, "2028-01": 90000}))
D["tech"] = ("Technologia: hosting, CDN, SaaS, AI API, CRM [zł]", steps({"2026-11": 4000, "2026-12": 9000, "2027-06": 12000, "2028-01": 15000}))
D["datalic"] = ("Licencje danych rynkowych (TGE i in.) [zł]", steps({"2026-12": 4000, "2027-07": 6000}))
D["research"] = ("Badania (Barometr Odbiorców: własna ankieta online) [zł]", steps({"2026-11": 0}))
D["legal"] = ("Księgowość, opłaty urzędowe, ubezpieczenia (dokumenty prawne przygotowuje rola Prawnik) [zł]", steps({"2026-11": 9000}))
D["office"] = ("Biuro/coworking, sprzęt, inne [zł]", steps({"2026-11": 6000, "2027-06": 9000}))
D["capex"] = ("Nakłady jednorazowe: znak towarowy i domeny (identyfikację, serwis i asystenta AI robimy sami) [zł]", [0.0] * N)
for l, v in {"2026-11": 4000}.items():
    D["capex"][1][idx[l]] = v
D["lead_qa_b2b"] = ("Koszt weryfikacji 1 leada MŚP (telefon, KRS, PPE) [zł]", steps({"2026-12": 25}))
D["lead_qa_hh"] = ("Koszt weryfikacji 1 leada GD/OZE (SMS/e-mail/telefon) [zł]", steps({"2027-01": 5}))
D["comm"] = ("Prowizja handlowa od reklamy/sponsoringu", steps({"2026-11": 0.10}))

# ---------- SCENARIUSZE ----------
SCEN = {"Ostrożny": (0.6, 0.85), "Bazowy": (1.0, 1.0), "Ambitny": (1.4, 1.15)}  # (mnożnik ruchu i konwersji -> przychód, mnożnik kosztów marketingu/zespołu)

def compute(rev_mult=1.0, cost_mult=1.0):
    g = lambda k, i: D[k][1][i]
    rows = {k: [0.0] * N for k in ["visits", "lead_b2b", "lead_hh", "lead_oze", "r_b2b", "r_big", "r_hh", "r_oze", "r_prog", "r_art", "r_pack", "r_news", "r_pro", "r_data", "r_events", "rev", "c_var", "c_fix", "c_capex", "cost", "ebitda", "cum"]}
    cum = 0
    for i in range(N):
        v = g("visits", i) * rev_mult
        lb = v * g("b2b_rate", i); lh = v * g("hh_rate", i); lo = v * g("oze_rate", i)
        r = dict(
            r_b2b=lb * g("b2b_sell", i) * g("b2b_price", i),
            r_big=g("big_deals", i) * g("big_fee", i) * rev_mult,
            r_hh=lh * g("hh_sell", i) * g("hh_price", i),
            r_oze=lo * g("oze_sell", i) * g("oze_price", i),
            r_prog=v / 1000 * g("rpm", i),
            r_art=g("spons_art", i) * g("spons_art_price", i) * rev_mult,
            r_pack=g("spons_pack", i) * g("spons_pack_price", i) * rev_mult,
            r_news=g("newsletter", i) * rev_mult,
            r_pro=g("pro_subs", i) * g("pro_price", i) * rev_mult,
            r_data=g("data_rev", i) * rev_mult,
            r_events=g("events", i) * g("event_margin", i) * rev_mult)
        rev = sum(r.values())
        cvar = lb * g("lead_qa_b2b", i) + (lh + lo) * g("lead_qa_hh", i) + (r["r_art"] + r["r_pack"] + r["r_news"] + r["r_events"]) * g("comm", i)
        cfix = (g("payroll", i) + g("marketing", i)) * cost_mult + g("experts", i) + g("tech", i) + g("datalic", i) + g("research", i) + g("legal", i) + g("office", i)
        cap = g("capex", i)
        cost = cvar + cfix + cap; e = rev - cost; cum += e
        for k, val in dict(visits=v, lead_b2b=lb, lead_hh=lh, lead_oze=lo, rev=rev, c_var=cvar, c_fix=cfix, c_capex=cap, cost=cost, ebitda=e, cum=cum, **r).items():
            rows[k][i] = val
    return rows

def fmt(x): return f"{x:,.0f}".replace(",", " ")

if __name__ == "__main__":
    out_xlsx = sys.argv[1]
    res = {s: compute(*mm) for s, mm in SCEN.items()}
    base = res["Bazowy"]
    # --- raport tekstowy do biznesplanu ---
    sel = [l for l in lab if l <= "2027-06"] + ["2027-09", "2027-12", "2028-06", "2028-12"]
    names = [("visits", "Wizyty / mies."), ("lead_b2b", "Leady MŚP"), ("r_b2b", "Leady MŚP [zł]"), ("r_big", "Duzi odbiorcy [zł]"), ("r_hh", "Leady GD [zł]"), ("r_oze", "Leady OZE [zł]"), ("r_prog", "Programmatic [zł]"), ("r_art", "Native [zł]"), ("r_pack", "Sponsoring sekcji [zł]"), ("r_news", "Newsletter [zł]"), ("r_pro", "PRO [zł]"), ("r_data", "Dane/raporty [zł]"), ("r_events", "Wydarzenia [zł]"), ("rev", "PRZYCHODY"), ("cost", "KOSZTY (z CAPEX)"), ("ebitda", "WYNIK"), ("cum", "Wynik skumulowany")]
    print("| Pozycja | " + " | ".join(sel) + " |"); print("|---|" + "---:|" * len(sel))
    for k, n in names:
        print(f"| {n} | " + " | ".join(fmt(base[k][idx[l]]) for l in sel) + " |")
    print()
    for s, r in res.items():
        def tot(a, b, key): return sum(r[key][idx[a]:idx[b] + 1])
        mn = min(r["cum"]); mi = r["cum"].index(mn)
        be = next((lab[i] for i in range(N) if r["ebitda"][i] > 0 and all(x > 0 for x in r["ebitda"][i:])), "po 2028-12")
        print(f"{s}: przych. XI26-VI27 {fmt(tot('2026-11','2027-06','rev'))}, koszty {fmt(tot('2026-11','2027-06','cost'))}; "
              f"przych.2027 {fmt(tot('2027-01','2027-12','rev'))}, wynik 2027 {fmt(tot('2027-01','2027-12','ebitda'))}; "
              f"przych.2028 {fmt(tot('2028-01','2028-12','rev'))}, wynik 2028 {fmt(tot('2028-01','2028-12','ebitda'))}; "
              f"max zapotrzebowanie {fmt(-mn)} ({lab[mi]}); trwała rentowność od {be}; rev VI27 {fmt(r['rev'][idx['2027-06']])}, rev XII28 {fmt(r['rev'][-1])}")

    # --- XLSX z formułami ---
    wb = Workbook(); ws = wb.active; ws.title = "Założenia"
    bold = Font(bold=True); inp = PatternFill("solid", fgColor="FFF2CC"); hdr = PatternFill("solid", fgColor="DDEBF7")
    ws["A1"] = "Energypedia.pl – model finansowy (scenariusz bazowy). Żółte komórki = założenia do edycji. Kwoty netto w zł."; ws["A1"].font = bold
    ws["A2"] = "Mnożnik przychodów (ruch, wolumeny)"; ws["B2"] = 1.0; ws["B2"].fill = inp
    ws["A3"] = "Mnożnik kosztów zespołu i marketingu"; ws["B3"] = 1.0; ws["B3"].fill = inp
    ws["D2"] = "Scenariusze: Ostrożny 0,6 / 0,85; Bazowy 1,0 / 1,0; Ambitny 1,4 / 1,15"
    R0 = 5
    ws.cell(R0, 1, "Założenie").font = bold
    for j, l in enumerate(lab):
        c = ws.cell(R0, 2 + j, l); c.font = bold; c.fill = hdr
    rowof = {}
    for n, (k, (name, vals)) in enumerate(D.items()):
        r = R0 + 1 + n; rowof[k] = r; ws.cell(r, 1, name)
        for j, v in enumerate(vals):
            c = ws.cell(r, 2 + j, v); c.fill = inp
            if k.endswith("rate"): c.number_format = "0.00%"
            elif k in ("b2b_sell", "hh_sell", "oze_sell", "comm"): c.number_format = "0%"
            else: c.number_format = "#,##0"
    ws.column_dimensions["A"].width = 62
    for j in range(N): ws.column_dimensions[get_column_letter(2 + j)].width = 12
    ws.freeze_panes = "B6"

    wm = wb.create_sheet("Model"); wm.cell(1, 1, "Rachunek wyników i przepływy (formuły – nie edytować)").font = bold
    for j, l in enumerate(lab):
        c = wm.cell(3, 2 + j, l); c.font = bold; c.fill = hdr
    A = lambda k, col: f"Założenia!{col}{rowof[k]}"
    lines = [
        ("visits", "Wizyty / mies.", lambda c: f"={A('visits',c)}*Założenia!$B$2"),
        ("lead_b2b", "Leady MŚP (kwalifikowane)", lambda c: f"={c}{{visits}}*{A('b2b_rate',c)}"),
        ("lead_hh", "Leady gospodarstw domowych", lambda c: f"={c}{{visits}}*{A('hh_rate',c)}"),
        ("lead_oze", "Leady OZE / efektywność", lambda c: f"={c}{{visits}}*{A('oze_rate',c)}"),
        (None, "PRZYCHODY", None),
        ("r_b2b", "Leady MŚP", lambda c: f"={c}{{lead_b2b}}*{A('b2b_sell',c)}*{A('b2b_price',c)}"),
        ("r_big", "Success fee – duzi odbiorcy", lambda c: f"={A('big_deals',c)}*{A('big_fee',c)}*Założenia!$B$2"),
        ("r_hh", "Leady GD", lambda c: f"={c}{{lead_hh}}*{A('hh_sell',c)}*{A('hh_price',c)}"),
        ("r_oze", "Leady OZE", lambda c: f"={c}{{lead_oze}}*{A('oze_sell',c)}*{A('oze_price',c)}"),
        ("r_prog", "Reklama programmatic", lambda c: f"={c}{{visits}}/1000*{A('rpm',c)}"),
        ("r_art", "Artykuły sponsorowane/native", lambda c: f"={A('spons_art',c)}*{A('spons_art_price',c)}*Założenia!$B$2"),
        ("r_pack", "Sponsoring sekcji/akademii", lambda c: f"={A('spons_pack',c)}*{A('spons_pack_price',c)}*Założenia!$B$2"),
        ("r_news", "Sponsoring newslettera", lambda c: f"={A('newsletter',c)}*Założenia!$B$2"),
        ("r_pro", "Abonament PRO", lambda c: f"={A('pro_subs',c)}*{A('pro_price',c)}*Założenia!$B$2"),
        ("r_data", "Dane, raporty, API", lambda c: f"={A('data_rev',c)}*Założenia!$B$2"),
        ("r_events", "Wydarzenia regionalne", lambda c: f"={A('events',c)}*{A('event_margin',c)}*Założenia!$B$2"),
        ("rev", "PRZYCHODY RAZEM", lambda c: f"=SUM({c}{{r_b2b}}:{c}{{r_events}})"),
        (None, "KOSZTY", None),
        ("c_var", "Koszty zmienne (weryfikacja leadów, prowizje)", lambda c: f"={c}{{lead_b2b}}*{A('lead_qa_b2b',c)}+({c}{{lead_hh}}+{c}{{lead_oze}})*{A('lead_qa_hh',c)}+({c}{{r_art}}+{c}{{r_pack}}+{c}{{r_news}}+{c}{{r_events}})*{A('comm',c)}"),
        ("c_fix", "Koszty stałe (zespół, marketing, tech, dane, badania, prawo, biuro)", lambda c: f"=({A('payroll',c)}+{A('marketing',c)})*Założenia!$B$3+{A('experts',c)}+{A('tech',c)}+{A('datalic',c)}+{A('research',c)}+{A('legal',c)}+{A('office',c)}"),
        ("c_capex", "Nakłady jednorazowe", lambda c: f"={A('capex',c)}"),
        ("cost", "KOSZTY RAZEM", lambda c: f"={c}{{c_var}}+{c}{{c_fix}}+{c}{{c_capex}}"),
        ("ebitda", "WYNIK GOTÓWKOWY (przed podatkiem)", lambda c: f"={c}{{rev}}-{c}{{cost}}"),
        ("cum", "WYNIK SKUMULOWANY (zapotrzebowanie na kapitał)", None),
    ]
    rr = {}; r = 4
    for k, _, _ in lines:
        if k: rr[k] = r
        r += 1
    r = 4
    for k, name, f in lines:
        wm.cell(r, 1, name).font = bold if (k is None or k in ("rev", "cost", "ebitda", "cum")) else Font()
        if k:
            for j in range(N):
                c = get_column_letter(2 + j)
                if k == "cum":
                    formula = f"={c}{rr['ebitda']}" if j == 0 else f"={get_column_letter(1 + j)}{r}+{c}{rr['ebitda']}"
                else:
                    formula = f(c).format(**rr)
                cell = wm.cell(r, 2 + j, formula); cell.number_format = "#,##0"
        r += 1
    wm.column_dimensions["A"].width = 58
    for j in range(N): wm.column_dimensions[get_column_letter(2 + j)].width = 12
    wm.freeze_panes = "B4"
    # podsumowania roczne
    r += 1; wm.cell(r, 1, "Podsumowanie").font = bold
    periods = [("XI 2026 – VI 2027", "2026-11", "2027-06"), ("Rok 2027", "2027-01", "2027-12"), ("Rok 2028", "2028-01", "2028-12")]
    for j, (pn, _, _) in enumerate(periods): wm.cell(r, 2 + j, pn).font = bold
    for k in ("rev", "cost", "ebitda"):
        r += 1; wm.cell(r, 1, dict(rev="Przychody", cost="Koszty", ebitda="Wynik")[k])
        for j, (_, a, b) in enumerate(periods):
            ca, cb = get_column_letter(2 + idx[a]), get_column_letter(2 + idx[b])
            wm.cell(r, 2 + j, f"=SUM({ca}{rr[k]}:{cb}{rr[k]})").number_format = "#,##0"
    r += 1; wm.cell(r, 1, "Maksymalne zapotrzebowanie na kapitał")
    wm.cell(r, 2, f"=-MIN(B{rr['cum']}:{get_column_letter(1 + N)}{rr['cum']})").number_format = "#,##0"

    # arkusz scenariuszy (wartości)
    wsc = wb.create_sheet("Scenariusze")
    wsc.append(["Wartości policzone skryptem dla trzech scenariuszy (statyczne)"]); wsc.append([])
    wsc.append(["Scenariusz", "Przychody XI26–VI27", "Koszty XI26–VI27", "Przychody 2027", "Wynik 2027", "Przychody 2028", "Wynik 2028", "Maks. zapotrzebowanie na kapitał", "Przychody VI 2027 (mies.)", "Przychody XII 2028 (mies.)"])
    for s, rres in res.items():
        t = lambda a, b, k: round(sum(rres[k][idx[a]:idx[b] + 1]))
        wsc.append([s, t("2026-11", "2027-06", "rev"), t("2026-11", "2027-06", "cost"), t("2027-01", "2027-12", "rev"), t("2027-01", "2027-12", "ebitda"), t("2028-01", "2028-12", "rev"), t("2028-01", "2028-12", "ebitda"), round(-min(rres["cum"])), round(rres["rev"][idx["2027-06"]]), round(rres["rev"][-1])])
    for row in wsc.iter_rows(min_row=4):
        for c in row[1:]: c.number_format = "#,##0"
    wsc.column_dimensions["A"].width = 14
    for col in "BCDEFGHIJ": wsc.column_dimensions[col].width = 20
    wb.save(out_xlsx)
