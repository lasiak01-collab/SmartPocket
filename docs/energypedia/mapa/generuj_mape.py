# Generuje arkusz Energypedia_mapa_tematow.xlsx z danych w mapa_dane.py (i opcjonalnie z dowodów popularności w dowody.py).
import os, re, sys, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from mapa_dane import OBSZARY, P, HASLA, PYTANIA, REGIONY
try:
    from dowody import DOWODY  # {fraza lub temat (małe litery): 'dowód ze źródłem'}
except ImportError:
    DOWODY = {}

OUT = os.path.join(os.path.dirname(__file__), 'Energypedia_mapa_tematow.xlsx')
NAVY, VOLT = '0B1B33', '00C2A8'
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor=NAVY)
WRAP = Alignment(wrap_text=True, vertical='top')
OBN = {k: n for k, n, *_ in OBSZARY}
FILAR = {k: u for k, _, u, *_ in OBSZARY}

def slug(t):
    t = t.lower().replace('ł', 'l')
    t = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode()
    t = re.sub(r'\(.*?\)', '', t)
    return re.sub(r'[^a-z0-9]+', '-', t).strip('-')

def norm(t):
    return re.sub(r'\s*\(.*?\)', '', t.lower()).strip()

def sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = H; c.fill = HF; c.alignment = WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws

def dowod(*frazy):
    out = []
    for f in frazy:
        for k, v in DOWODY.items():
            if k in f.lower() and v not in out:
                out.append(v)
    return ' | '.join(out) if out else '[DO WERYFIKACJI w narzędziu SEO]'

SCHEMA = {'przewodnik': 'Article + BreadcrumbList + FAQPage', 'tracker': 'NewsArticle (aktualizowany) + BreadcrumbList', 'artykuł': 'Article + BreadcrumbList + FAQPage',
          'analiza': 'Article + Dataset + BreadcrumbList', 'poradnik': 'Article (HowTo jako struktura treści) + BreadcrumbList + FAQPage', 'wzór pisma': 'Article + DigitalDocument + BreadcrumbList', 'katalog': 'CollectionPage + ItemList + BreadcrumbList', 'profil sprzedawcy': 'WebPage (about: Organization) + BreadcrumbList', 'profil operatora': 'WebPage (about: Organization) + BreadcrumbList'}

wb = Workbook()
ws = wb.active; ws.title = 'Instrukcja'
info = [
    ['Energypedia.pl – mapa tematów (wersja 1.0, 21.10.2026 r.)'],
    [''],
    ['Co jest w arkuszu'],
    ['Obszary – 10 obszarów tematycznych z adresem strony przewodniej (filaru) i frazą główną.'],
    ['Publikacje MVP – %d publikacji do 31.12.2026 r.: 5 przewodników, 2 trackery, 20 artykułów i porównań, 5 analiz, 25 poradników, 5 wzorów pism, katalog: 5 sprzedawców z urzędu i 5 operatorów sieci ze stronami zbiorczymi.' % len(P)],
    ['Hasła – encyklopedia: lista priorytetowa; kolumna „MVP” wskazuje 120 haseł na start.'],
    ['Pytania long tail – pytania odbiorców przypisane do stron docelowych; to także zestaw do comiesięcznego pomiaru cytowań w wyszukiwarkach AI.'],
    ['Kanibalizacja – frazy, o które walczyłyby dwie strony, z decyzją, która strona jest główna.'],
    ['Regiony – 16 stron województw.'],
    ['Taksonomia – wymiary klasyfikacji treści.'],
    ['Technika SEO i GEO – wymagania techniczne całego serwisu.'],
    [''],
    ['Jak liczymy priorytet'],
    ['Wynik = popyt × wartość × szansa × pilność, każdy czynnik w skali 1–3 (maks. 81).'],
    ['Popyt to ocena jakościowa redakcji, dopóki nie wpiszemy danych z narzędzia SEO (Google Search Console, Senuto, Semstorm, Ahrefs lub Planer słów kluczowych Google Ads).'],
    ['Kolumna „Dowód popularności” zawiera tylko informacje ze źródłem. Gdzie nie ma źródła, jest znacznik [DO WERYFIKACJI w narzędziu SEO]. Nie podajemy wolumenów bez źródła.'],
    [''],
    ['Zasady nadrzędne: ZASADY_Energypedia.md oraz zespol/06_redaktor_seo_aeo_geo.md (anatomia artykułu, jedna intencja = jedna strona).'],
]
for r in info:
    ws.append(r)
ws['A1'].font = Font(bold=True, size=14); ws['A3'].font = Font(bold=True); ws['A13'].font = Font(bold=True)
ws.column_dimensions['A'].width = 140

sheet(wb, 'Obszary', ['Kod', 'Obszar', 'Strona przewodnia (URL)', 'Fraza główna filaru', 'Zakres', 'Publikacji MVP', 'Haseł', 'Pytań long tail'],
      [[k, n, u, f, o, sum(1 for p in P if p['obszar'] == k), len(HASLA.get(k, [])), len(PYTANIA.get(k, []))] for k, n, u, f, o in OBSZARY],
      [24, 34, 40, 26, 70, 14, 10, 14])

pub_frazy = {}
rows = []
for i, p in enumerate(sorted(P, key=lambda p: -(p['popyt'] * p['wartosc'] * p['szansa'] * p['pilnosc'])), 1):
    w = p['popyt'] * p['wartosc'] * p['szansa'] * p['pilnosc']
    pub_frazy[norm(p['fraza'])] = p['url']
    rows.append([i, p['typ'], OBN[p['obszar']], p['h1'], p['fraza'], p['poboczne'], p['intencja'], p['segment'], p['url'], FILAR[p['obszar']],
                 p['format_pz'], SCHEMA[p['typ']], p['popyt'], p['wartosc'], p['szansa'], p['pilnosc'], w, p['sezon'],
                 {1: 'Fala 1 (do 20.11)', 2: 'Fala 2 (do 4.12)', 'start': 'Tydzień startu'}[p['fala']], dowod(p['fraza'], p['poboczne'])])
sheet(wb, 'Publikacje MVP', ['Lp.', 'Typ', 'Obszar', 'Tytuł roboczy (H1)', 'Fraza główna', 'Frazy poboczne i long tail', 'Intencja', 'Segment', 'URL', 'Link do filaru',
                             'Format pozycji zerowej', 'Dane strukturalne', 'Popyt (1–3)', 'Wartość (1–3)', 'Szansa (1–3)', 'Pilność (1–3)', 'Priorytet (maks. 81)',
                             'Sezonowość', 'Termin', 'Dowód popularności (ze źródłem)'],
      rows, [5, 11, 22, 48, 26, 50, 18, 14, 40, 30, 12, 30, 9, 9, 9, 9, 11, 22, 18, 60])

hrows, kanib = [], []
for o, lista in HASLA.items():
    for j, h in enumerate(lista):
        url = '/wiedza/' + slug(h)
        mvp = 'tak' if j < round(120 * len(lista) / sum(len(v) for v in HASLA.values())) else ''
        glowna = pub_frazy.get(norm(h))
        if glowna:
            kanib.append([h, url, glowna, 'Fraza ogólna należy do publikacji. Hasło celuje w definicję („co to jest …”) i linkuje do publikacji jako strony głównej.'])
        hrows.append([OBN[o], h, url, 'co to jest ' + norm(h), 'DefinedTerm + BreadcrumbList', mvp, glowna or '', FILAR[o]])
sheet(wb, 'Hasła', ['Obszar', 'Hasło', 'URL', 'Fraza docelowa hasła', 'Dane strukturalne', 'MVP (120)', 'Strona główna dla frazy ogólnej (jeśli inna)', 'Link do filaru'],
      hrows, [30, 42, 48, 46, 26, 10, 40, 34])

prow = []
for o, lista in PYTANIA.items():
    for q in lista:
        prow.append([OBN[o], q, FILAR[o], 'akapit 40–60 słów pod nagłówkiem-pytaniem', '[DO WERYFIKACJI w narzędziu SEO]', ''])
sheet(wb, 'Pytania long tail', ['Obszar', 'Pytanie', 'Strona docelowa (do doprecyzowania przy planowaniu)', 'Format odpowiedzi', 'Popyt', 'Cytowana przez AI? (pomiar miesięczny)'],
      prow, [30, 70, 40, 34, 28, 26])

sheet(wb, 'Kanibalizacja', ['Fraza', 'Hasło (URL)', 'Strona główna dla frazy', 'Decyzja'], kanib, [34, 44, 44, 90])

sheet(wb, 'Regiony', ['Województwo', 'URL', 'Fraza główna', 'Zawartość', 'Operatorzy sieci na terenie województwa'],
      [[r, '/regiony/' + slug(r), 'prąd ' + r + ' operator', 'operator sieci i jego stawki, przerwy i awarie, dotacje wojewódzkie, kontakt do instytucji', '[DO WERYFIKACJI z map obszarów działania OSD]'] for r in REGIONY],
      [22, 36, 34, 70, 44])

sheet(wb, 'Taksonomia', ['Wymiar', 'Wartości', 'Zastosowanie'], [
    ['Obszar (10)', ', '.join(k for k, *_ in OBSZARY), 'nawigacja, filary, okruszki, targetowanie reklam'],
    ['Segment (4)', 'dom, mikro, msp, duze-samorzady', 'personalizacja, sekcja „Co to znaczy dla Ciebie?”, CTA'],
    ['Typ treści', 'news, analiza, opinia, przewodnik, poradnik, porównanie, lista, hasło, tracker, dane, wzór, partner', 'szablony, dane strukturalne, raporty'],
    ['Region', '16 województw, operatorzy sieci', 'strony regionalne, treści lokalne'],
    ['Tagi kontrolowane', 'słownik pojęć = hasła encyklopedii (arkusz Hasła)', 'linkowanie wewnętrzne, strony tematów /tematy/{tag}'],
    ['Encje', 'instytucje (URE, PSE, TGE, ARE, UOKiK), sprzedawcy energii, operatorzy sieci, akty prawne, programy', 'dane strukturalne, linkowanie, GEO'],
    ['Etap 2 (od II kw. 2027 r.)', 'dział gazu (/gaz/), rankingi (/rankingi/), zestawienia produktów (/zestawienia/)', 'rozbudowa taksonomii o obszar „gaz” i typy „ranking”, „zestawienie”'],
], [22, 90, 60])

TECH = [
    ('Indeksacja', 'Mapy witryny: ogólna, newsy (dla Google News), obrazy; zgłoszenie w Google Search Console i Bing Webmaster Tools; IndexNow przy każdej publikacji i aktualizacji.'),
    ('Indeksacja', 'Treść renderowana po stronie serwera (bez treści dostępnej tylko przez JavaScript); kod 200 dla stron, 301 przy łączeniu treści, 410 dla usuniętych.'),
    ('Indeksacja', 'Parametry filtrów i wyszukiwanie wewnętrzne: noindex i adres kanoniczny do wersji bazowej; paginacja z linkami do kolejnych stron.'),
    ('Adresy', 'Krótkie, po polsku, bez polskich znaków, małe litery, łączone dywizem; bez roku w adresach treści evergreen (rok tylko w danych, np. /dane/barometr-rynku/2026-12).'),
    ('Adresy', 'Jedna intencja = jeden adres; przed publikacją sprawdzenie arkusza Kanibalizacja i istniejących stron.'),
    ('Dane strukturalne', 'NewsMediaOrganization w całym serwisie (z polami publishingPrinciples, correctionsPolicy, ownershipFundingInfo, masthead); Article lub NewsArticle, BreadcrumbList, Person (autorzy), DefinedTerm (hasła), Dataset (dane), FAQPage tam, gdzie jest sekcja FAQ.'),
    ('Dane strukturalne', 'Walidacja w Rich Results Test i walidatorze schema.org przy każdym nowym szablonie.'),
    ('Meta', 'Unikalny title (do 60 znaków) i description (do 155 znaków); robots max-image-preview:large; Open Graph i karty social z obrazem 1200×630; data publikacji i aktualizacji widoczna na stronie i w danych strukturalnych.'),
    ('Wydajność', 'Core Web Vitals w zakresie „dobry” dla 75. percentyla: LCP do 2,5 s, INP do 200 ms, CLS do 0,1; obrazy AVIF/WebP z wymiarami; zarezerwowana wysokość slotów reklamowych.'),
    ('Dostępność', 'WCAG 2.2 AA w szablonach: semantyczny HTML, nagłówki w kolejności, teksty alternatywne, kontrast, obsługa klawiaturą.'),
    ('Linkowanie wewnętrzne', 'Każda publikacja linkuje do filaru obszaru i do haseł użytych pojęć; filar linkuje do wszystkich publikacji obszaru; raport stron osieroconych co miesiąc.'),
    ('Linkowanie wewnętrzne', 'Okruszki na każdej stronie; sekcja „Czytaj dalej” sterowana relacjami w CMS.'),
    ('Treść (AEO)', 'Anatomia artykułu z zespol/06_redaktor_seo_aeo_geo.md: lead z odpowiedzią, blok „W skrócie” (komponent KeyFacts), nagłówki-pytania, tabele HTML, listy kroków, FAQ, źródła, dziennik zmian.'),
    ('GEO', 'Fragmenty samodzielne pod każdym H2; liczby z jednostką, datą i źródłem w tym samym zdaniu; własne dane (Barometr); cytaty ekspertów z imieniem i funkcją.'),
    ('GEO', 'Spójne encje: pełne nazwy instytucji przy pierwszym użyciu; strona O nas i Organization z sameAs; docelowo wpis w Wikidata.'),
    ('GEO', 'robots.txt: dopuszczone boty wyszukiwarek i wyszukiwarek AI; decyzja o botach trenujących modele należy do naczelnego; llms.txt jako uzupełnienie (wpływ niepotwierdzony).'),
    ('GEO', 'Pomiar: arkusz Pytania long tail jako stały zestaw pytań sprawdzanych co miesiąc w ChatGPT, Perplexity, Gemini, Copilot i AI Overviews.'),
    ('Preferowane źródła Google', 'Wezwanie „Dodaj do preferowanych źródeł” (link https://www.google.com/preferences/source?q=energypedia.pl lub oficjalny przycisk z Search Central) pod artykułami, w panelu „Więcej” i w stopce; pomiar kliknięć; regularne newsy, bo funkcja działa w sekcji „Najważniejsze artykuły”. [DO WERYFIKACJI: wytyczne i kod przycisku]'),
    ('E-E-A-T', 'Strony: Jak pracujemy, Jak zarabiamy, Redakcja, Sprostowania, strony autorów z biogramem; podpis autora i recenzenta przy każdej publikacji.'),
    ('Google News i Discover', 'Zgłoszenie w Google Publisher Center; obrazy co najmniej 1200 px szerokości; tytuły bez clickbaitu; jasna data publikacji.'),
    ('Kanały', 'RSS/Atom: całość, obszary, trackery; API danych Barometru; wykresy do osadzenia z linkiem zwrotnym.'),
    ('Język', 'ZASADY_Energypedia.md: poprawna polszczyzna, bez długiego myślnika, typografia liczb i dat; lang="pl" w HTML.'),
]
sheet(wb, 'Technika SEO i GEO', ['Obszar', 'Wymaganie'], TECH, [24, 140])

wb.save(OUT)
print('Zapisano', OUT, '| publikacje:', len(P), '| hasła:', sum(len(v) for v in HASLA.values()), '| pytania:', sum(len(v) for v in PYTANIA.values()), '| kanibalizacja:', len(kanib))
